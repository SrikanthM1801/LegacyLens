"""SharePoint -> knowledge: text extraction, delta sync (add / unchanged / modify / delete / unsupported /
failed), the source registry, audit events, and the admin endpoints. Offline: mock Graph, memory stores."""

import asyncio
import io
from datetime import datetime, timedelta, timezone

import pytest
from fastapi.testclient import TestClient

from app.config import get_settings
from app.graph import get_graph
from app.ingestion.file_text import UnsupportedFile, extract_text
from app.ingestion.sharepoint import artifact_id_for, sync_sharepoint
from app.main import app
from app.services import audit
from app.services.consent_store import get_consent_store
from app.services.graph_api import GraphError, get_graph_api
from app.services.search import get_search
from app.services.source_registry import get_source_registry

SITE = "https://contoso.sharepoint.com/sites/LegacyLensProject"
DRIVE = "drive-1"


def docx_bytes(*paragraphs: str) -> bytes:
    import docx
    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        c.post("/ingest")
        yield c


@pytest.fixture(autouse=True)
def site(monkeypatch):
    monkeypatch.setattr(get_settings(), "sharepoint_site_url", SITE)
    graph = get_graph_api()
    graph.reset()
    graph.add_site(SITE, "site-1", "LegacyLensProject", DRIVE)
    get_source_registry().reset()
    get_consent_store().clear()
    audit.get_audit_log().clear()
    return graph


def sync(full=False):
    return asyncio.run(sync_sharepoint(full=full))


def found(phrase: str, item_id: str) -> bool:
    return artifact_id_for(item_id) in {c.artifact_id for c in get_search().search(phrase, top_k=10)}


def events(action: str):
    now = datetime.now(timezone.utc)
    return [e for e in audit.get_audit_log().query(now - timedelta(hours=1), now + timedelta(minutes=1)) if e.action == action]


def test_extract_text_from_office_and_text_files():
    import openpyxl
    from pptx import Presentation

    assert "Quokka rollout" in extract_text("plan.docx", docx_bytes("Quokka rollout", "Step two"))
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Architecture overview"
    buf = io.BytesIO(); prs.save(buf)
    assert "Slide 1: Architecture overview" in extract_text("deck.pptx", buf.getvalue())
    wb = openpyxl.Workbook(); wb.active.append(["Owner", "Service"]); wb.active.append(["Kavi", "Billing"])
    buf = io.BytesIO(); wb.save(buf)
    assert "Kavi | Billing" in extract_text("owners.xlsx", buf.getvalue())
    assert extract_text("notes.md", b"# Title\r\n\r\n  body  ") == "# Title\nbody"
    with pytest.raises(UnsupportedFile):
        extract_text("archive.zip", b"PK")


def test_first_sync_loads_files_records_registry_consent_and_audit(site):
    site.put_file(DRIVE, "f1", "Release Runbook.docx", docx_bytes("Run the quokkaflip deploy script before 6pm"),
                  folder="/drive/root:/Runbooks")
    site.put_file(DRIVE, "f2", "notes.txt", b"The marmotsync job owner is Kaviraj")
    site.put_file(DRIVE, "f3", "logo.png", b"\x89PNG")
    report = sync()
    assert (report.added, report.unsupported, report.failed) == (2, 1, 0) and report.chunks_indexed >= 2
    assert found("quokkaflip deploy script", "f1") and found("marmotsync job owner", "f2")
    reg = {f.name: f for f in get_source_registry().list()}
    assert reg["Release Runbook.docx"].status == "loaded" and reg["Release Runbook.docx"].path == "/Runbooks"
    assert reg["Release Runbook.docx"].web_url.startswith("https://") and reg["logo.png"].status == "unsupported"
    assert get_graph().artifact_node(artifact_id_for("f1"))["type"] == "doc"  # counts as documentation
    assert get_consent_store().get("site-1").status == "granted"
    files = events("capture.file")
    assert {e.details["file"] for e in files} == {"Release Runbook.docx", "notes.txt", "logo.png"}
    assert all(e.details["change"] == "added" for e in files)
    assert "quokkaflip" not in str([e.model_dump() for e in files])  # no file contents in the audit trail
    assert len(events("sync.sharepoint")) == 1 and len(events("consent.granted")) == 1


def test_only_changes_are_processed_afterwards(site):
    site.put_file(DRIVE, "f1", "spec.txt", b"The wombatqueue retries three times")
    sync()
    audit.get_audit_log().clear()
    assert sync().added == 0 and events("capture.file") == []        # nothing new in the delta
    assert sync(full=True).unchanged == 1                              # full re-read skips by cTag


def test_modified_file_replaces_old_knowledge(site):
    site.put_file(DRIVE, "f1", "spec.txt", b"The wombatqueue retries three times")
    sync()
    site.put_file(DRIVE, "f1", "spec.txt", b"The pangolinbus now replaces the old queue", version=2)
    report = sync()
    assert report.modified == 1
    assert found("pangolinbus replaces", "f1") and not found("wombatqueue retries", "f1")
    assert events("capture.file")[-1].details["change"] == "modified"


def test_deleted_file_is_forgotten(site):
    site.put_file(DRIVE, "f9", "old.txt", b"The narwhalcache is warmed nightly")
    sync()
    site.delete_file(DRIVE, "f9")
    assert sync().deleted == 1
    assert not found("narwhalcache warmed nightly", "f9")
    assert not get_graph().has_node(artifact_id_for("f9"))
    assert get_source_registry().get("f9").status == "deleted"


def test_download_failure_is_recorded_not_raised(site, monkeypatch):
    site.put_file(DRIVE, "bad", "broken.docx", b"not really a docx")
    report = sync()
    assert report.failed == 1 and get_source_registry().get("bad").status == "failed"
    assert events("capture.file")[0].outcome == "failure"

    async def boom(*_a):
        raise GraphError(403, "resolve site")

    monkeypatch.setattr(site, "resolve_site", boom)
    assert sync().errors and events("sync.sharepoint")[-1].outcome == "failure"


def test_admin_endpoints(client, site):
    site.put_file(DRIVE, "f1", "a.txt", b"The ibexindex is rebuilt weekly")
    r = client.post("/knowledge/sync")
    assert r.status_code == 200 and r.json()["added"] == 1
    body = client.get("/knowledge/sources").json()
    assert body["total"] == 1 and body["by_status"] == {"loaded": 1} and body["last_sync"]
    assert client.get("/knowledge/sources", params={"status": "failed"}).json()["total"] == 0
    assert client.post("/knowledge/rebuild").status_code == 409  # sample data is on in tests


def test_graph_remove_artifact_keeps_shared_evidence():
    from app.graph.networkx_store import NetworkXGraphStore
    from app.models import Edge, Entity, EntityType, RelationType

    g = NetworkXGraphStore()
    g.add_entity(Entity(id="t1", name="Topic", type=EntityType.SYSTEM, artifact_ids=["a1", "a2"]))
    g.add_entity(Entity(id="t2", name="Only a1", type=EntityType.SYSTEM, artifact_ids=["a1"]))
    g.add_entity(Entity(id="p1", name="Pat", type=EntityType.PERSON))
    for a in ("a1", "a2"):
        g.add_artifact_node(a, "doc", a)
    g.add_edge(Edge(source="p1", target="t1", relation=RelationType.AUTHORED, weight=2, evidence_artifact_ids=["a1", "a2"]))
    g.add_edge(Edge(source="a1", target="t1", relation=RelationType.MENTIONS, evidence_artifact_ids=["a1"]))
    assert g.remove_artifact("a1") and not g.remove_artifact("a1")
    [kept] = g.edges(source="p1")
    assert kept.evidence_artifact_ids == ["a2"] and kept.weight == 1
    assert g.get_entity("t2") is None and g.get_entity("t1").artifact_ids == ["a2"]


def test_indexing_failure_marks_files_failed_and_retries_next_run(site, monkeypatch):
    import app.ingestion.sharepoint as sp

    site.put_file(DRIVE, "f1", "a.txt", b"The okapiqueue drains hourly")
    real = sp._index
    monkeypatch.setattr(sp, "_index", lambda *a: (_ for _ in ()).throw(RuntimeError("429")))
    report = sync()
    assert report.failed == 1 and get_source_registry().get("f1").status == "failed"
    assert get_source_registry().delta_link(DRIVE) is None  # not advanced
    monkeypatch.setattr(sp, "_index", real)
    assert sync().added == 1 and found("okapiqueue drains hourly", "f1")
