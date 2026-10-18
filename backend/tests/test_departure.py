"""Departure prevention, offline on the sample project (Ravi holds the payment gateway knowledge):
detection (out-of-office / away / owner), alert, 1:1 interview, approval -> risk drops, handover pack."""

import asyncio
import io
from datetime import datetime, timedelta, timezone

import pytest
from fastapi.testclient import TestClient

from app.bot import agent, departure
from app.config import get_settings
from app.main import app
from app.models import ConsentRecord
from app.services import audit, interview_store
from app.services.consent_store import get_consent_store
from app.services.docx_export import markdown_to_docx
from app.services.graph_api import get_graph_api
from app.services.teams_connector import get_connector
from tests.test_consent import BOT, CHANNEL, TEAM

RAVI_AAD = "aad-ravi"
OWNER = {"id": "29:owner", "name": "Owner Olga", "aadObjectId": "aad-owner"}
NOW = datetime(2026, 10, 18, 9, tzinfo=timezone.utc)


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        c.post("/ingest")
        yield c


@pytest.fixture(autouse=True)
def fresh(client, monkeypatch):
    client.post("/ingest")  # every test starts from the planted risks
    for s in (departure.departures(), departure.links(), get_consent_store(), audit.get_audit_log()):
        s.clear()
    interview_store.clear()
    get_graph_api().reset()
    get_connector().reset()
    get_graph_api().owners["group-1"] = {"aad-owner"}
    get_graph_api().users["ravi@contoso.com"] = (RAVI_AAD, "Ravi Kumar")
    get_consent_store().save(ConsentRecord(scope_id=CHANNEL["id"], scope_type="channel", status="granted"))
    monkeypatch.setattr(get_settings(), "sharepoint_site_url", "https://contoso.sharepoint.com/sites/p")
    yield


def actions():
    now = datetime.now(timezone.utc)
    return [e.action for e in audit.get_audit_log().query(now - timedelta(hours=1), now + timedelta(minutes=1))]


def test_parse_until_and_out_of_office_windows():
    assert departure.parse_until("away until 3 Nov", NOW).date().isoformat() == "2026-11-03"
    assert departure.parse_until("away till 2027-01-05", NOW).date().isoformat() == "2027-01-05"
    assert departure.parse_until("away until 2 Jan", NOW).year == 2027      # next occurrence
    assert departure.parse_until("away for a while", NOW) is None
    sched = lambda s, e: {"status": "scheduled", "scheduledStartDateTime": {"dateTime": s, "timeZone": "UTC"},
                          "scheduledEndDateTime": {"dateTime": e, "timeZone": "UTC"}}
    assert departure.ooo_window(sched("2026-10-20T00:00:00", "2026-11-03T00:00:00"), NOW, 14)[0].day == 20
    assert departure.ooo_window(sched("2026-12-20T00:00:00", "2027-01-03T00:00:00"), NOW, 14) is None  # too far
    assert departure.ooo_window(sched("2026-10-01T00:00:00", "2026-10-05T00:00:00"), NOW, 14) is None  # over
    assert departure.ooo_window({"status": "alwaysEnabled"}, NOW, 14)[0] == NOW
    assert departure.ooo_window({"status": "disabled"}, NOW, 14) is None and departure.ooo_window(None, NOW, 14) is None


def test_link_maps_a_teams_user_to_the_document_persona():
    link, reply = departure.link_person(RAVI_AAD, "R K (Teams)", "Ravi Kumar", "Owner Olga")
    assert "Linked" in reply and "ravi" in link.person_ids
    assert "ravi" in departure.person_ids_for(RAVI_AAD, "R K (Teams)")
    assert departure.link_person(RAVI_AAD, "x", "Nobody Here", "o")[0] is None
    assert "ravi" in departure.person_ids_for("aad-other", "Ravi Kumar")  # same name matches without a link


def test_full_handover_alert_interview_approve_risk_drops_and_pack_saved():
    departure.link_person(RAVI_AAD, "Ravi Kumar", "Ravi Kumar", "Owner Olga")
    d, reply = asyncio.run(departure.start(aad_id=RAVI_AAD, name="Ravi Kumar", source="out_of_office",
                                           starts_at=NOW + timedelta(days=2), ends_at=NOW + timedelta(days=16)))
    assert "depend mostly on Ravi" in reply and d.status == "interviewing"
    assert d.at_risk and d.at_risk[0].topic == "Payment Gateway" and d.at_risk[0].share >= 0.5
    sent = get_connector().sent
    assert sent[0][0] == d.conversation_id and "before you're away" in str(sent[0][1])          # intro
    assert "Question 1 of" in str(sent[1][1])                                                     # first question
    assert sent[-1][0] == CHANNEL["id"] and "Ravi Kumar is away" in str(sent[-1][1])             # channel alert
    draft = [p for p in get_graph_api().uploads if p.startswith("LegacyLens Handover/")]
    assert draft and "handover pack" in str(sent[-1][1]).lower()

    total = sum(len(interview_store.get(s).questions) for s in d.session_ids)
    for i in range(total):
        card, text = departure.answer(departure.active_for(RAVI_AAD), None, None,
                                      "Check `gwctl pool-status` first, then restart pg-adapter; escalate to the acquirer NOC after 10 minutes.")
    assert "check your answers" in text and "Approve & publish" in str(card)
    before = d.at_risk[0].risk_score
    active = departure.active_for(RAVI_AAD)
    assert departure.begin_approval(active) is None and departure.begin_approval(active).startswith("⏳")  # double click
    asyncio.run(departure.approve_in_background(d.id))
    assert "Published" in get_connector().sent[-1][1]["text"]  # confirmation in the 1:1 chat
    final = departure.departures().get(d.id)
    gw = final.at_risk[0]
    assert final.status == "completed" and gw.risk_after is not None and gw.risk_after < before
    assert any("Knowledge captured from Ravi" in str(a) for _, a in get_connector().sent)
    import docx
    doc = docx.Document(io.BytesIO(get_graph_api().uploads[draft[0]]))
    assert any("Ravi" in p.text for p in doc.paragraphs)
    acts = actions()
    for a in ("departure.detected", "departure.interview_started", "departure.alert_posted", "departure.answered",
              "departure.knowledge_published", "departure.handover_published"):
        assert a in acts, a


def test_nobody_at_risk_gets_an_all_clear():
    d, reply = asyncio.run(departure.start(aad_id="aad-new", name="New Joiner", source="self"))
    assert d.status == "no_risk" and "nothing is at risk" in reply
    assert "No topic depends mostly on New" in str(get_connector().sent[-1][1])


def test_out_of_office_detection_starts_once():
    departure.link_person(RAVI_AAD, "Ravi Kumar", "Ravi Kumar", "o")
    get_graph_api().ooo[RAVI_AAD] = {"status": "alwaysEnabled"}
    started = asyncio.run(departure.check_out_of_office([RAVI_AAD, "aad-guest"], {RAVI_AAD: "Ravi Kumar"}))
    assert len(started) == 1 and started[0].source == "out_of_office"
    assert asyncio.run(departure.check_out_of_office([RAVI_AAD], {RAVI_AAD: "Ravi Kumar"})) == []  # already handled


def _say(client, text, frm, conversation=CHANNEL, value=None):
    mention = [{"type": "mention", "mentioned": BOT, "text": "<at>LegacyLens</at>"}]
    body = {"type": "message", "id": "m", "channelId": "msteams", "serviceUrl": "https://smba.trafficmanager.net/in/",
            "from": frm, "recipient": BOT, "conversation": conversation, "deliveryMode": "expectReplies",
            "channelData": {"team": {**TEAM, "aadGroupId": "group-1"}} if conversation is CHANNEL else {},
            "text": f"<at>LegacyLens</at> {text}" if conversation is CHANNEL else text,
            "entities": mention if conversation is CHANNEL else []}
    if value is not None:
        body["value"] = value
    r = client.post("/api/messages", json=body)
    assert r.status_code == 200, r.text
    return [a for a in r.json()["activities"] if a.get("type") != "typing"]


def test_teams_commands_owner_handover_link_and_away(client, monkeypatch):
    kavi = {"id": "29:kavi", "name": "Kavi", "aadObjectId": "aad-kavi"}
    assert "Only a team owner" in _say(client, "handover ravi@contoso.com", kavi)[0]["text"]
    assert "Linked" in _say(client, "link ravi@contoso.com as Ravi Kumar", OWNER)[0]["text"]
    reply = _say(client, "handover ravi@contoso.com until 3 Nov", OWNER)[0]["text"]
    assert "depend mostly on Ravi" in reply
    d = departure.active_for(RAVI_AAD)
    assert d.source == "owner" and d.ends_at.month == 11
    # the interview continues in Ravi's 1:1 chat
    ravi = {"id": "29:ravi", "name": "Ravi Kumar", "aadObjectId": RAVI_AAD}
    personal = {"id": d.conversation_id, "conversationType": "personal"}
    out = _say(client, "Check gwctl pool-status first and restart pg-adapter", ravi, personal)
    assert out[0]["text"] in ("Thanks! Next one:", "That's all. Please check your answers:")
    assert "already helping" in _say(client, "away until 30 Nov", ravi)[0]["text"]


def test_markdown_to_docx_handles_headings_lists_links():
    import docx
    data = markdown_to_docx("# Title\n## Part\n- **Bold** item [1]\n1. step\n> note\nSee [doc](https://x.y/z)")
    d = docx.Document(io.BytesIO(data))
    texts = [p.text for p in d.paragraphs]
    assert texts[0] == "Title" and "Bold item [1]" in texts and "See doc (https://x.y/z)" in texts


def test_live_extraction_links_people_named_in_documents_and_chat_authors():
    from app.ingestion.extractor import LLMExtraction, LLMExtractor, RuleExtractor
    from app.models import Artifact, ArtifactType, Person

    class FakeLLM:
        model_name = "gpt-test"

        def chat_json(self, system, user, schema):
            return LLMExtraction.model_validate({"entities": [
                {"name": "RAID log", "type": "process", "owners": ["Dinesh Udhyakumar", "Product Owner"]},
                {"name": "Use AKS", "type": "decision", "decided_by": "Subash", "owners": []}]})

    now = datetime.now(timezone.utc)
    doc = Artifact(type=ArtifactType.DOC, source_id="d", author="person_up", timestamp=now, title="Governance",
                   body="...", metadata={"live": True})
    chat = Artifact(type=ArtifactType.CHAT, source_id="c", author="person_aad_kavi", timestamp=now, title="chat",
                    body="...", metadata={"live": True})
    res = LLMExtractor(FakeLLM(), rules=RuleExtractor(catalog=False)).extract(
        [doc, chat], [Person(id="person_up", name="Uploader"), Person(id="person_aad_kavi", name="A, Kaviraj")])
    names = {e.name for e in res.entities if e.type == "person"}
    assert {"Dinesh Udhyakumar", "Subash"} <= names and "Product Owner" not in names
    rel = {(e.source, e.relation, e.target) for e in res.edges}
    assert ("person_doc_dinesh_udhyakumar", "PARTICIPATED_IN", "proc_raid_log") in rel
    assert ("person_doc_subash", "DECIDED", "dec_use_aks") in rel
    assert ("person_aad_kavi", "DISCUSSED", "proc_raid_log") in rel
    assert ("proc_raid_log", "DOCUMENTED_BY", doc.id) in rel


def test_approve_click_is_acknowledged_immediately_and_published_in_background(client, monkeypatch):
    spawned = []
    monkeypatch.setattr(agent, "_spawn", spawned.append)
    departure.link_person(RAVI_AAD, "Ravi Kumar", "Ravi Kumar", "o")
    d, _ = asyncio.run(departure.start(aad_id=RAVI_AAD, name="Ravi Kumar", source="self"))
    total = sum(len(interview_store.get(s).questions) for s in d.session_ids)
    for _ in range(total):
        departure.answer(departure.active_for(RAVI_AAD), None, None, "Run gwctl pool-status, restart pg-adapter, escalate after 10 min.")
    ravi = {"id": "29:ravi", "name": "Ravi Kumar", "aadObjectId": RAVI_AAD}
    personal = {"id": d.conversation_id, "conversationType": "personal"}
    [ack] = _say(client, "", ravi, personal, value={"verb": departure.APPROVE_VERB, "departure_id": d.id})
    assert ack["text"].startswith("⏳ Publishing") and len(spawned) == 1
    [again] = _say(client, "", ravi, personal, value={"verb": departure.APPROVE_VERB, "departure_id": d.id})
    assert "Already publishing" in again["text"] and len(spawned) == 1
    asyncio.run(spawned.pop())
    assert departure.departures().get(d.id).status == "completed"
