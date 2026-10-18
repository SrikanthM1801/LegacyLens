"""
Markdown (as produced by agents/generator.py) -> .docx bytes, so handover packs land in SharePoint
as normal Word documents. Handles headings, bullet / numbered lists, block quotes, **bold** and
[links](url); citation markers like [3] are kept as text.
"""

from __future__ import annotations

import io
import re

_INLINE = re.compile(r"\*\*(.+?)\*\*|\[([^\]]+)\]\((https?://[^)]+)\)")


def _runs(paragraph, text: str) -> None:
    pos = 0
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        if m.group(1) is not None:
            paragraph.add_run(m.group(1)).bold = True
        else:
            paragraph.add_run(f"{m.group(2)} ({m.group(3)})")
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def markdown_to_docx(markdown: str, title: str = "") -> bytes:
    import docx

    doc = docx.Document()
    if title:
        doc.core_properties.title = title
    for raw in markdown.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if m := re.match(r"^(#{1,4})\s+(.*)", line):
            doc.add_heading(m.group(2).strip("* "), level=min(len(m.group(1)), 3))
        elif m := re.match(r"^\s*[-*]\s+(.*)", line):
            _runs(doc.add_paragraph(style="List Bullet"), m.group(1))
        elif m := re.match(r"^\s*\d+[.)]\s+(.*)", line):
            _runs(doc.add_paragraph(style="List Number"), m.group(1))
        elif line.startswith(">"):
            p = doc.add_paragraph()
            _runs(p, line.lstrip("> "))
            for r in p.runs:
                r.italic = True
        else:
            _runs(doc.add_paragraph(), line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
